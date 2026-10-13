#!/usr/bin/env python3
"""Cut a finished run's results down to a few samples, for teaching.

A training session needs output to read, which normally means running the
pipeline first. When a full batch has already been processed, this is the same
thing for free: take that run's workbook and mapped tables, keep only the rows
for the samples the trainees were given, and write the result somewhere they can
point `OUT` at. No scheduler, no service units.

Filenames are preserved, so every command in docs/tutorial-first-run-gadi.md
works unchanged against the output directory:

    ls $OUT/*.xlsx
    cat $OUT/*_review_required.tsv

Usage:
  ./scripts/subset_results_for_training.py \\
      --results /scratch/rg42/AGAR/intermediates/2024/B07 \\
      --samples-from /scratch/rg42/training/2026-10-12 \\
      --out /scratch/rg42/training/2026-10-12/reference_results

  ./scripts/subset_results_for_training.py \\
      --results /path/to/finished/run \\
      --samples 24GNB-1752,24GNB-1745,24GNB-1775 \\
      --out /path/to/output

Exits non-zero if a requested sample is absent from the source run, since a
silently short reference is worse than none.
"""

from __future__ import annotations

import argparse
import csv
import shutil
import sys
from pathlib import Path

# Header spellings that identify the sample column. Compared lowercased with
# runs of whitespace collapsed, matching read_metadata_sample_rows() in
# validate_metadata_samples.py.
SAMPLE_HEADERS = {
    "sample name",
    "sample",
    "sample_name",
    "samplename",
    "name",
    "isolate",
    "strain",
    "tag",
    "query",
}

# Tables that are one row per run rather than per sample; copied whole.
COPY_WHOLE = {"software versions", "versions", "citations"}

# Scratch and scheduler noise. Nextflow work trees in particular are enormous and
# hold a copy of everything, which would both dwarf the real output and expose
# every other sample in the run.
SKIP_DIRS = {"_work", "work", "tmp", "pipeline_logs", ".nextflow", "nf-reports"}
SKIP_SUFFIXES = {".log", ".err", ".out"}

# PBS delivers job logs as <name>.o<jobid> / <name>.e<jobid>.
def is_scheduler_log(name: str) -> bool:
    for marker in (".o", ".e"):
        head, _, tail = name.rpartition(marker)
        if head and tail.isdigit():
            return True
    return False


def norm(value: object) -> str:
    return " ".join(str(value if value is not None else "").split()).lower()


def find_sample_column(header: list[object]) -> int | None:
    """Index of the sample column, or None if this table is not per-sample."""
    for index, cell in enumerate(header):
        if norm(cell) in SAMPLE_HEADERS:
            return index
    return None


def collect_samples_from_sheets(root: Path) -> list[str]:
    """Sample ids from every *_samplesheet.txt under root, in file order."""
    samples: list[str] = []
    seen: set[str] = set()
    sheets = sorted(root.rglob("*_samplesheet.txt"))
    if not sheets:
        raise SystemExit(f"No *_samplesheet.txt found under {root}")

    for sheet in sheets:
        with sheet.open(newline="", encoding="utf-8-sig") as handle:
            first = handle.readline()
            delimiter = "\t" if "\t" in first else ","
            handle.seek(0)
            rows = list(csv.reader(handle, delimiter=delimiter))
        if not rows:
            continue
        # Row 1 is always a header here, matching how the pipeline reads these.
        for row in rows[1:]:
            if not row:
                continue
            sample = row[0].strip()
            if sample and sample not in seen:
                seen.add(sample)
                samples.append(sample)
    return samples


def filter_delimited(src: Path, wanted: set[str]) -> tuple[str, list[list[str]], set[str]]:
    """Read a TSV/CSV and keep the header plus rows for `wanted`.

    Returns (delimiter, rows-including-header, sample ids seen). Nothing is
    written here: the caller decides, so that a table with no matching rows does
    not leave an empty directory behind.
    """
    with src.open(newline="", encoding="utf-8-sig") as handle:
        first = handle.readline()
        delimiter = "\t" if "\t" in first else ","
        handle.seek(0)
        rows = list(csv.reader(handle, delimiter=delimiter))

    if not rows:
        return delimiter, [], set()

    header = rows[0]
    column = find_sample_column(header)
    if column is None:
        column = 0  # these tables lead with the sample id even when unlabelled

    kept = [header]
    seen: set[str] = set()
    for row in rows[1:]:
        if len(row) <= column:
            continue
        sample = row[column].strip()
        if sample in wanted:
            kept.append(row)
            seen.add(sample)
    return delimiter, kept, seen


def write_delimited(dest: Path, delimiter: str, rows: list[list[str]]) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle, delimiter=delimiter, lineterminator="\n").writerows(rows)


def subset_delimited(src: Path, dest: Path, wanted: set[str]) -> tuple[int, set[str]]:
    """Filter one TSV/CSV into dest. Returns (rows kept, sample ids seen)."""
    delimiter, rows, seen = filter_delimited(src, wanted)
    write_delimited(dest, delimiter, rows)
    return max(len(rows) - 1, 0), seen


def subset_workbook(src: Path, dest: Path, wanted: set[str]) -> tuple[dict[str, int], set[str]]:
    """Filter every per-sample sheet of a workbook. Returns (per-sheet counts, ids seen)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit(
            "openpyxl is required to subset the workbook.\n"
            "On Gadi use the MLST env's interpreter, which already has it:\n"
            "  /g/data/rg42/bactopia_datasets/envs/mlst_env/bin/python3 "
            "scripts/subset_results_for_training.py ..."
        )

    workbook = load_workbook(src)
    counts: dict[str, int] = {}
    seen: set[str] = set()

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            continue
        if norm(worksheet.title) in COPY_WHOLE:
            counts[worksheet.title] = max(len(rows) - 1, 0)
            continue

        header = list(rows[0])
        column = find_sample_column(header)
        if column is None:
            # Not a per-sample table; leave it as it is.
            counts[worksheet.title] = max(len(rows) - 1, 0)
            continue

        keep = [rows[0]]
        for row in rows[1:]:
            if len(row) <= column:
                continue
            value = row[column]
            sample = str(value).strip() if value is not None else ""
            if sample in wanted:
                keep.append(row)
                seen.add(sample)

        worksheet.delete_rows(1, worksheet.max_row)
        for row in keep:
            worksheet.append(row)
        counts[worksheet.title] = len(keep) - 1

    workbook.save(dest)
    return counts, seen


def human(size: int) -> str:
    value = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.0f}{unit}" if unit == "B" else f"{value:.1f}{unit}"
        value /= 1024
    return f"{value:.1f}GB"


def all_samples_in(table: Path) -> set[str]:
    """Every sample id in a mapped results table, so skips can be explained."""
    try:
        with table.open(newline="", encoding="utf-8-sig") as handle:
            first = handle.readline()
            delimiter = "\t" if "\t" in first else ","
            handle.seek(0)
            rows = list(csv.reader(handle, delimiter=delimiter))
    except (OSError, UnicodeDecodeError):
        return set()
    if not rows:
        return set()
    column = find_sample_column(rows[0])
    if column is None:
        column = 0
    return {
        row[column].strip()
        for row in rows[1:]
        if len(row) > column and row[column].strip()
    }


def subset_outputs(results: Path, out: Path, wanted: set[str], others: set[str],
                   dry_run: bool) -> tuple[int, int, list[str], dict[str, int]]:
    """Copy the per-sample output tree for the chosen samples.

    Three kinds of file, handled differently:
      - anything whose path names one of our samples: copied (assemblies,
        per-sample tool output, ST131Typer directories)
      - merged tables with a sample column: filtered to our samples
      - everything else: skipped and reported, because a merged file we cannot
        filter would carry every other sample in the run into the training set
    """
    copied = 0
    total = 0
    skipped: list[str] = []
    reasons: dict[str, int] = {}

    def skip(rel, why: str) -> None:
        skipped.append(f"{rel} ({why})")
        reasons[why] = reasons.get(why, 0) + 1

    top_level = sorted(
        p for p in results.iterdir()
        if p.is_dir() and p.name not in SKIP_DIRS and not p.name.startswith(".")
    )

    for directory in top_level:
        for src in sorted(directory.rglob("*")):
            if not src.is_file():
                continue
            rel = src.relative_to(results)
            if any(part in SKIP_DIRS or part.startswith(".") for part in rel.parts[:-1]):
                continue
            if src.suffix.lower() in SKIP_SUFFIXES or is_scheduler_log(src.name):
                continue

            text = str(rel)
            dest = out / rel

            if any(sample in text for sample in wanted):
                if not dry_run:
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(src, dest)
                copied += 1
                total += src.stat().st_size
                continue

            if src.suffix.lower() in {".tsv", ".csv"}:
                try:
                    with src.open(newline="", encoding="utf-8-sig") as handle:
                        first = handle.readline()
                    if not first.strip():
                        skip(rel, "empty")
                        continue
                    delimiter = "\t" if "\t" in first else ","
                    header = next(csv.reader([first], delimiter=delimiter))
                except (OSError, UnicodeDecodeError, StopIteration):
                    skip(rel, "unreadable")
                    continue

                if find_sample_column(header) is None:
                    skip(rel, "merged table, no sample column")
                    continue

                # Filter first, write only if anything survived. Writing then
                # deleting would leave an empty directory per other sample, and
                # a tree full of empty sample directories reads as though every
                # sample is present.
                delimiter, rows, _ = filter_delimited(src, wanted)
                if len(rows) <= 1:
                    skip(rel, "no matching rows")
                    continue
                if not dry_run:
                    write_delimited(dest, delimiter, rows)
                    total += dest.stat().st_size
                copied += 1
                continue

            if any(other in text for other in others):
                skip(rel, "belongs to another sample")
            else:
                skip(rel, "not per-sample")

    # Belt and braces: anything that still ended up empty goes, deepest first,
    # so a directory emptied by this loop is removed in the same pass.
    if not dry_run and out.is_dir():
        for path in sorted(out.rglob("*"), key=lambda p: len(p.parts), reverse=True):
            if path.is_dir():
                try:
                    path.rmdir()
                except OSError:
                    pass  # not empty, which is the normal case

    return copied, total, skipped, reasons


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Subset a finished run's results to a few samples, for training.",
    )
    parser.add_argument("--results", required=True, type=Path,
                        help="RESULTS_ROOT of the finished run to take results from")
    parser.add_argument("--out", required=True, type=Path,
                        help="Directory to write the subset into (created if absent)")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--samples",
                       help="Comma- or space-separated sample ids")
    group.add_argument("--samples-from", type=Path,
                       help="Directory to read *_samplesheet.txt files from, "
                            "e.g. a training tree")
    parser.add_argument("--no-outputs", action="store_true",
                        help="Only the workbook and mapped tables. By default the "
                             "per-sample output tree is included too (assemblies, "
                             "tools, Kleborate, FimTyper, ST131Typer, consolidated), "
                             "so the result looks like a real run")
    parser.add_argument("--list-skipped", action="store_true",
                        help="Name every file left out of the output tree")
    parser.add_argument("--dry-run", action="store_true",
                        help="Report what would be written, change nothing")
    args = parser.parse_args()

    if not args.results.is_dir():
        raise SystemExit(f"No such results directory: {args.results}")

    if args.samples:
        samples = [s for s in args.samples.replace(",", " ").split() if s]
    else:
        samples = collect_samples_from_sheets(args.samples_from)
    if not samples:
        raise SystemExit("No sample ids given")
    wanted = set(samples)

    print(f"Source run : {args.results}")
    print(f"Samples    : {len(samples)} -> {', '.join(samples)}")
    print(f"Destination: {args.out}")
    print()

    # The files the tutorial tells people to open, plus the review table.
    patterns = [
        "*_samplesheet_with_results_mlst_reviewed.tsv",
        "*_samplesheet_with_results.tsv",
        "*_samplesheet_with_results_review_required.tsv",
        "*_results.xlsx",
    ]
    targets: list[Path] = []
    for pattern in patterns:
        targets.extend(sorted(p for p in args.results.glob(pattern) if p.is_file()))
    if not targets:
        raise SystemExit(
            f"Found no results files under {args.results}.\n"
            "Expected at least one of: " + ", ".join(patterns)
        )

    if not args.dry_run:
        args.out.mkdir(parents=True, exist_ok=True)

    found: set[str] = set()
    for src in targets:
        dest = args.out / src.name
        if args.dry_run:
            print(f"  [dry-run] {src.name}")
            continue
        if src.suffix.lower() == ".xlsx":
            counts, seen = subset_workbook(src, dest, wanted)
            found |= seen
            print(f"  {src.name}")
            for sheet, n in counts.items():
                print(f"      {sheet:<34} {n:>5} rows")
        else:
            n, seen = subset_delimited(src, dest, wanted)
            found |= seen
            print(f"  {src.name:<52} {n:>5} rows")

    if not args.no_outputs:
        run_samples: set[str] = set()
        for target in targets:
            if target.suffix.lower() == ".tsv":
                run_samples |= all_samples_in(target)
        others = run_samples - wanted

        print()
        print("Per-sample output tree:")
        copied, total, skipped, reasons = subset_outputs(
            args.results, args.out, wanted, others, args.dry_run)
        print(f"  kept {copied} file(s){'' if args.dry_run else ', ' + human(total)}")
        for why, count in sorted(reasons.items(), key=lambda kv: -kv[1]):
            print(f"  left out {count:>5}  {why}")
        if skipped and args.list_skipped:
            for entry in skipped:
                print(f"      {entry}")
        elif skipped:
            print("      (--list-skipped names them)")

    if args.dry_run:
        print("\nDry run - nothing written.")
        return 0

    missing = sorted(wanted - found)
    print()
    if missing:
        print(f"WARNING: {len(missing)} sample(s) not found in the source run:")
        for sample in missing:
            print(f"  - {sample}")
        print("\nThe subset is incomplete. Check these samples were part of that run.")
        return 1

    print(f"All {len(samples)} samples found. Point users' OUT at:")
    print(f"  {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
