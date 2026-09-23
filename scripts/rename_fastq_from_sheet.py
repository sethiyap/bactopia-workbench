#!/usr/bin/env python3
"""Rename FASTQ files to their AGAR sample ids using an isolate -> AGAR sheet.

Legacy AGAR data predates the AGRF naming convention: the files carry a local
isolate id (19-004-0109_R1.fq.gz) rather than an AGAR accession, so they fail
^[0-9]{2}GNB-[0-9]+R?$. In AGAR mode submit_workbench_pipeline.sh filters the
FOFN on that pattern, and normalize_agar_fastq_sample_names.sh exits 0 while
reporting them INVALID -- so the samples are dropped with nothing failing. This
renames them from the sheet that maps each isolate id to its AGAR id.

Nothing is derived from the filename. An isolate with no sheet entry is reported
and left alone, because an AGAR accession cannot be guessed from a local id.

AGRF-layout files belong to normalize_agar_fastq_sample_names.sh and are skipped.
"""

# str | None annotations would need 3.10; the deployment runs older pythons.
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

AGAR_SAMPLE_REGEX = re.compile(r"^[0-9]{2}GNB-[0-9]+R?$")
# <sample>_R1.fastq.gz -- the layout legacy deliveries use.
PAIR_REGEX = re.compile(r"^(?P<sample>.+)_R(?P<read>[12])\.(?P<ext>fastq|fq)\.gz$")
# <sample>_<flowcell>_<barcode>_L001_R1.fastq.gz -- an AGRF delivery.
AGRF_REGEX = re.compile(r"^.+_[^_]+_[^_]+_L[0-9]{3}_R[12]\.(fastq|fq)\.gz$")

# Tried in order when the columns are not named explicitly. The isolate side is
# what the files are currently called; the AGAR side is what they should become.
FROM_CANDIDATES = (
    "isolate id", "isolate", "isolate name", "legacy id", "old id",
    "original id", "lab id", "lab number", "local id", "sample id",
    "old sample name",
)
# "sample" is last-resort: in an AGAR sheet the Sample column holds the AGAR id
# and Isolate holds the local one, but a sheet using "sample" for the isolate
# would resolve the wrong way round, so the explicit AGAR titles win first.
TO_CANDIDATES = (
    "agar id", "agar", "agar number", "agar sample", "agar sample name",
    "agar accession", "new id", "new sample name", "sample name", "sample",
)


def normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).strip()).lower()


def read_rows(path: Path) -> list[list[str]]:
    """Every non-blank row of the sheet as strings, header row included."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit(
                f"Reading {path.name} needs openpyxl, which usually lives in "
                "MLST_ENV rather than the system python. Run this with "
                "$MLST_ENV/bin/python3, or export the sheet to .tsv and pass that."
            )
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in workbook.active.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
    else:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            first_line = handle.readline()
            if not first_line:
                raise SystemExit(f"Sheet is empty: {path}")
            # Same sniff as validate_metadata_samples.py: a tab anywhere in the
            # header line means tab-delimited, otherwise comma.
            delimiter = "\t" if "\t" in first_line else ","
            handle.seek(0)
            rows = list(csv.reader(handle, delimiter=delimiter))

    return [row for row in rows if any(str(cell).strip() for cell in row)]


def resolve_column(headers: list[str], wanted: str | None, candidates, role: str,
                   flag: str) -> int:
    available = ", ".join(repr(h) for h in headers if h)
    if wanted:
        target = normalize_header(wanted)
        if target in headers:
            return headers.index(target)
        raise SystemExit(f"No column named {wanted!r}.\nAvailable columns: {available}")
    for candidate in candidates:
        if candidate in headers:
            return headers.index(candidate)
    raise SystemExit(
        f"Could not work out which column holds the {role}.\n"
        f"Available columns: {available}\n"
        f"Name it explicitly with {flag}."
    )


def build_mapping(path: Path, from_col: str | None, to_col: str | None,
                  allow_nonstandard: bool) -> dict[str, str]:
    """isolate id -> AGAR id, refusing any sheet that cannot be applied safely."""
    rows = read_rows(path)
    if not rows:
        raise SystemExit(f"Sheet has no usable rows: {path}")

    headers = [normalize_header(cell) for cell in rows[0]]
    from_idx = resolve_column(headers, from_col, FROM_CANDIDATES, "isolate id", "--from-col")
    to_idx = resolve_column(headers, to_col, TO_CANDIDATES, "agar id", "--to-col")
    if from_idx == to_idx:
        raise SystemExit("The isolate-id and agar-id columns must be different.")

    mapping: dict[str, str] = {}
    reverse: dict[str, str] = {}
    conflicts: list[str] = []
    collisions: list[str] = []
    malformed: list[str] = []

    for line_number, row in enumerate(rows[1:], start=2):
        if len(row) <= max(from_idx, to_idx):
            continue
        old = str(row[from_idx]).strip()
        new = str(row[to_idx]).strip()
        if not old or not new:
            continue

        # One isolate with two AGAR ids is unresolvable; two isolates sharing one
        # AGAR id would merge distinct samples into a single assembly.
        if old in mapping and mapping[old] != new:
            conflicts.append(f"  line {line_number}: {old} -> {new} (already {mapping[old]})")
            continue
        if new in reverse and reverse[new] != old:
            collisions.append(f"  line {line_number}: {old} and {reverse[new]} both -> {new}")
            continue
        if not AGAR_SAMPLE_REGEX.match(new):
            malformed.append(f"  line {line_number}: {old} -> {new}")

        mapping[old] = new
        reverse[new] = old

    problems = []
    if conflicts:
        problems.append("One isolate id maps to more than one AGAR id:\n" + "\n".join(conflicts))
    if collisions:
        problems.append("Two isolate ids map to the same AGAR id:\n" + "\n".join(collisions))
    if malformed and not allow_nonstandard:
        problems.append(
            "These AGAR ids do not match ^[0-9]{2}GNB-[0-9]+R?$ and would be "
            "dropped by the AGAR FOFN filter:\n" + "\n".join(malformed)
            + "\nFix the sheet, or pass --allow-nonstandard."
        )
    if problems:
        raise SystemExit("\n\n".join(problems))
    if not mapping:
        raise SystemExit(f"No usable id pairs found in: {path}")

    print(
        f"Resolved {len(mapping)} id pairs from {path.name} "
        f"[{headers[from_idx]!r} -> {headers[to_idx]!r}]"
    )
    return mapping


def plan(fastq_dir: Path, mapping: dict[str, str]) -> tuple[list[tuple[Path, Path, str, str]], dict[str, int]]:
    """Work out every rename before any of them happens.

    Building the whole plan first means a conflict found late cannot leave the
    directory half-converted.
    """
    renames: list[tuple[Path, Path, str, str]] = []
    counts = {"checked": 0, "already": 0, "skipped": 0}
    seen_mates: dict[str, set[str]] = {}

    for path in sorted(fastq_dir.iterdir()):
        if not path.is_file() or not path.name.endswith((".fastq.gz", ".fq.gz")):
            continue
        counts["checked"] += 1

        if AGRF_REGEX.match(path.name):
            print(f"SKIP     {path.name}\n         AGRF layout; use normalize_agar_fastq_sample_names.sh")
            counts["skipped"] += 1
            continue

        matched = PAIR_REGEX.match(path.name)
        if not matched:
            print(f"SKIP     {path.name}\n         not a recognised paired-end FASTQ name")
            counts["skipped"] += 1
            continue

        sample = matched.group("sample")
        if AGAR_SAMPLE_REGEX.match(sample):
            counts["already"] += 1
            continue

        if sample not in mapping:
            print(f"SKIP     {path.name}\n         '{sample}' has no entry in the sheet")
            counts["skipped"] += 1
            continue

        new_sample = mapping[sample]
        target = path.with_name(f"{new_sample}_R{matched.group('read')}.{matched.group('ext')}.gz")
        if target.exists() and target != path:
            raise SystemExit(f"Refusing to rename: target already exists: {target.name}")

        seen_mates.setdefault(sample, set()).add(matched.group("read"))
        renames.append((path, target, sample, new_sample))

    # The FOFN builder fails on a missing R2, so catch it here rather than at submission.
    for sample, reads in sorted(seen_mates.items()):
        if reads != {"1", "2"}:
            missing = ({"1", "2"} - reads).pop()
            raise SystemExit(f"Refusing to rename: {sample} has no R{missing}")

    return renames, counts


def list_columns_mode() -> int:
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("sheet", type=Path)
    parser.add_argument("--list-columns", action="store_true")
    args, _ = parser.parse_known_args()
    if not args.sheet.is_file():
        raise SystemExit(f"Sheet not found: {args.sheet}")
    for header in (normalize_header(cell) for cell in read_rows(args.sheet)[0]):
        if header:
            print(header)
    return 0


def main() -> int:
    if "--list-columns" in sys.argv[1:]:
        return list_columns_mode()

    parser = argparse.ArgumentParser(
        description="Rename FASTQ files to their AGAR sample ids from a mapping sheet.",
        epilog="Dry run by default; pass --apply to perform the renames.",
    )
    parser.add_argument("sheet", type=Path, help="Isolate-id -> AGAR-id sheet (.tsv/.csv/.txt/.xlsx)")
    parser.add_argument("fastq_dir", type=Path, help="Directory of FASTQ files")
    parser.add_argument("--apply", action="store_true", help="Perform the renames")
    parser.add_argument("--from-col", help="Column holding the current isolate id")
    parser.add_argument("--to-col", help="Column holding the AGAR id")
    parser.add_argument("--list-columns", action="store_true", help="Print the sheet's columns and exit")
    parser.add_argument("--allow-nonstandard", action="store_true",
                        help="Accept AGAR ids that break ^[0-9]{2}GNB-[0-9]+R?$")
    parser.add_argument("--map-file", type=Path,
                        help="Where to write the old->new map (default: <fastq_dir>/fastq_rename_map.tsv)")
    args = parser.parse_args()

    if not args.sheet.is_file():
        raise SystemExit(f"Sheet not found: {args.sheet}")

    if not args.fastq_dir.is_dir():
        raise SystemExit(f"Directory not found: {args.fastq_dir}")

    mapping = build_mapping(args.sheet, args.from_col, args.to_col, args.allow_nonstandard)
    renames, counts = plan(args.fastq_dir, mapping)

    map_file = args.map_file or args.fastq_dir / "fastq_rename_map.tsv"
    status = "RENAMED" if args.apply else "PLANNED"
    print()
    with map_file.open("w", encoding="utf-8") as handle:
        handle.write("status\tsheet\toriginal_filename\tfinal_filename\toriginal_sample\tfinal_sample\n")
        for source, target, old_sample, new_sample in renames:
            if args.apply:
                source.rename(target)
            handle.write(
                f"{status}\t{args.sheet.name}\t{source.name}\t{target.name}\t{old_sample}\t{new_sample}\n"
            )
            print(f"RENAME   {source.name} -> {target.name}")

    print(
        f"\nSheet:      {args.sheet.name}"
        f"\nChecked:    {counts['checked']}"
        f"\n{'Renamed:   ' if args.apply else 'To rename: '} {len(renames)}"
        f"\nAlready ok: {counts['already']}"
        f"\nSkipped:    {counts['skipped']}"
        f"\nMap file:   {map_file}"
    )
    if not args.apply and renames:
        print("\nDry run only. Re-run with --apply to perform these renames.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
